"""
Genera el archivo Excel a partir de los movimientos extraídos de cualquier banco.
"""

from typing import Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from bancos.base import Movimiento, InfoCuenta

COLOR_DEPOSITO = "E8F5E9"
COLOR_RETIRO   = "FFEBEE"
COLOR_SALDO    = "E3F2FD"
COLOR_ALT_ROW  = "FAFAFA"
COLOR_WHITE    = "FFFFFF"
COLOR_TOTAL_D  = "C8E6C9"
COLOR_TOTAL_R  = "FFCDD2"

MONEY_FMT = '#,##0.00'

THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)


def _cell(ws, row, col, value=None, *, bold=False, color=None, bg=None,
          align="left", fmt=None, wrap=False, size=9, fg="212121"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, size=size, color=fg)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border = THIN
    if bg:
        c.fill = PatternFill("solid", fgColor=bg)
    if fmt:
        c.number_format = fmt
    return c


def export_to_excel(
    movimientos: list[Movimiento],
    output_path: str,
    info: InfoCuenta,
    empresa_nombre: str = "",
    usuario: str = "",
) -> str:
    """
    Genera el Excel de movimientos.

    Parametros adicionales:
      empresa_nombre : nombre de la empresa obtenido del catalogo por RFC
      usuario        : nombre del usuario que realizo la conversion
    """
    color_header = info.banco and _get_bank_color(info.banco) or "2C3E50"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movimientos"

    # ── Fila 1: Empresa (si se encontro) o Titulo del banco ───────────────────
    ws.merge_cells("A1:G1")
    c1 = ws["A1"]
    if empresa_nombre:
        c1.value = empresa_nombre.upper()
    else:
        c1.value = f"Estado de Cuenta {info.banco}"
    c1.font = Font(bold=True, size=14, color="FFFFFF")
    c1.fill = PatternFill("solid", fgColor=color_header)
    c1.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # ── Fila 2: Banco + Periodo ────────────────────────────────────────────────
    ws.merge_cells("A2:G2")
    c2 = ws["A2"]
    banco_periodo = f"{info.banco}  —  {info.periodo or ''}"
    if empresa_nombre and info.rfc:
        banco_periodo = f"RFC: {info.rfc}   |   {banco_periodo}"
    c2.value = banco_periodo
    c2.font = Font(bold=True, size=10, color="FFFFFF")
    c2.fill = PatternFill("solid", fgColor=color_header)
    c2.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # ── Fila 3: Info de cuenta + usuario ──────────────────────────────────────
    ws.merge_cells("A3:G3")
    c3 = ws["A3"]
    parts = []
    if info.numero_cuenta:
        parts.append(f"Cuenta: {info.numero_cuenta}")
    if info.titular:
        parts.append(f"Titular: {info.titular}")
    if not empresa_nombre and info.rfc:
        parts.append(f"RFC: {info.rfc}")
    parts.append(f"{len(movimientos)} movimientos")
    if usuario:
        parts.append(f"Generado por: {usuario}")
    c3.value = "   |   ".join(parts)
    c3.font = Font(italic=True, size=8, color="555555")
    c3.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 14

    row_offset = 1   # filas extra por fila 2 adicional

    # ── Encabezados ───────────────────────────────────────────────────────────
    HDR_ROW = 4
    # Col A = ID ESTABLE (llave de sync con la BD). Las demas se corren +1.
    headers = ["ID", "FECHA", "FOLIO", "DESCRIPCION", "DEPOSITO", "RETIRO", "SALDO"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=HDR_ROW, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=color_header)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = THIN
    ws.row_dimensions[HDR_ROW].height = 22

    # ── Saldo inicial ─────────────────────────────────────────────────────────
    data_start = HDR_ROW + 1
    if info.saldo_inicial is not None:
        ws.merge_cells(f"B{data_start}:D{data_start}")
        c = ws.cell(row=data_start, column=2, value="SALDO INICIAL DEL PERIODO ANTERIOR")
        c.font = Font(bold=True, italic=True, size=9, color="444444")
        c.alignment = Alignment(horizontal="right", vertical="center")
        c.border = THIN
        for col in (3, 4):
            ws.cell(row=data_start, column=col).border = THIN
        _cell(ws, data_start, 7, info.saldo_inicial,
              fmt=MONEY_FMT, bg=COLOR_SALDO, bold=True, align="right", size=9)
        for col in (5, 6):
            ws.cell(row=data_start, column=col).border = THIN
        data_start += 1

    # ── Filas de datos ────────────────────────────────────────────────────────
    for i, mov in enumerate(movimientos):
        row = data_start + i
        alt = COLOR_ALT_ROW if i % 2 else COLOR_WHITE

        _cell(ws, row, 1, getattr(mov, "id_mov", None) or None, bg=alt, align="center")
        _cell(ws, row, 2, mov.fecha, bg=alt, align="center")
        _cell(ws, row, 3, mov.folio, bg=alt, align="center")
        desc = (mov.descripcion or "").replace(" | ", "\n")
        _cell(ws, row, 4, desc, bg=alt, wrap=True)
        num_lineas = desc.count("\n") + 1
        ws.row_dimensions[row].height = max(18, num_lineas * 14)
        _cell(ws, row, 5, mov.deposito, fmt=MONEY_FMT, bg=COLOR_DEPOSITO if mov.deposito else alt, align="right")
        _cell(ws, row, 6, mov.retiro,   fmt=MONEY_FMT, bg=COLOR_RETIRO   if mov.retiro   else alt, align="right")
        _cell(ws, row, 7, mov.saldo,    fmt=MONEY_FMT, bg=COLOR_SALDO,   align="right")

    # ── Totales ────────────────────────────────────────────────────────────────
    total_row = data_start + len(movimientos)
    ws.merge_cells(f"B{total_row}:D{total_row}")
    c = ws.cell(row=total_row, column=2, value="TOTALES DEL PERIODO")
    c.font = Font(bold=True, size=10, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=color_header)
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.border = THIN
    for col in (1, 3, 4):
        ws.cell(row=total_row, column=col).border = THIN
        ws.cell(row=total_row, column=col).fill = PatternFill("solid", fgColor=color_header)

    total_dep = info.total_depositos or sum(m.deposito or 0 for m in movimientos)
    total_ret = info.total_retiros   or sum(m.retiro   or 0 for m in movimientos)
    _cell(ws, total_row, 5, total_dep, fmt=MONEY_FMT, bg=COLOR_TOTAL_D, bold=True, align="right")
    _cell(ws, total_row, 6, total_ret, fmt=MONEY_FMT, bg=COLOR_TOTAL_R, bold=True, align="right")
    saldo_fin = info.saldo_final or (movimientos[-1].saldo if movimientos else None)
    _cell(ws, total_row, 7, saldo_fin, fmt=MONEY_FMT, bg=COLOR_SALDO,   bold=True, align="right")

    # ── Anchos ────────────────────────────────────────────────────────────────
    for col, w in enumerate([8, 14, 12, 52, 16, 16, 18], 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = f"A{data_start}"
    ws.auto_filter.ref = f"A{HDR_ROW}:G{total_row - 1}"

    wb.save(output_path)
    return output_path


def _get_bank_color(banco: str) -> str:
    colors = {
        "Santander": "CC0000",
        "BBVA":      "004481",
        "Banregio":  "005A8E",
        "Banorte":   "006341",
        "Inbursa":   "E31837",
        "HSBC":      "DB0011",
        "Banamex":   "003DA5",
        "Scotiabank":"EC111A",
    }
    for key, color in colors.items():
        if key.lower() in banco.lower():
            return color
    return "2C3E50"
