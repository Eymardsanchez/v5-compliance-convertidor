"""
catalogo_cuentas.py — Catálogo de Cuentas Contables por empresa

Storage: %APPDATA%\\ConvertidorBancos\\cuentas_{RFC}.json
Una empresa → su propio archivo de catálogo.

Importación flexible desde Excel COI (detecta columnas por nombre).
Columnas esperadas: No. Cuenta / Cuenta, Nombre, Tipo (D/A), Nivel, Moneda.
"""

from __future__ import annotations

import json
import os
import unicodedata
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Optional


# ── Persistencia ──────────────────────────────────────────────────────────────

def _app_dir() -> Path:
    p = Path(os.environ.get("APPDATA", Path.home())) / "ConvertidorBancos"
    p.mkdir(parents=True, exist_ok=True)
    return p

def _path(rfc: str) -> Path:
    rfc_safe = rfc.strip().upper().replace(" ", "_")
    return _app_dir() / f"cuentas_{rfc_safe}.json"


# ── Dataclass ─────────────────────────────────────────────────────────────────

@dataclass
class CuentaContable:
    cuenta:  str   # "1110-001-000"
    nombre:  str   # "Bancos Nacionales"
    tipo:    str   # "D"=deudora / "A"=acreedora
    nivel:   int   # 1-5
    moneda:  str   # "MN" / "ME" / "BI"

    @property
    def label(self) -> str:
        return f"{self.cuenta}  {self.nombre}"


# ── Helpers de normalización ──────────────────────────────────────────────────

def _norm(s: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFKD", str(s).upper())
        if not unicodedata.combining(c)
    ).strip()


def _col_idx(hdrs: list[str], *palabras: str) -> Optional[int]:
    for i, h in enumerate(hdrs):
        hn = _norm(h)
        if any(p in hn for p in palabras):
            return i
    return None


# ── API pública ───────────────────────────────────────────────────────────────

def cargar_catalogo(rfc: str) -> list[CuentaContable]:
    """Carga el catálogo de cuentas de la empresa. Retorna [] si no existe."""
    try:
        raw = json.loads(_path(rfc).read_text(encoding="utf-8"))
        return [CuentaContable(**r) for r in raw]
    except Exception:
        return []


def guardar_catalogo(rfc: str, cuentas: list[CuentaContable]) -> None:
    _path(rfc).write_text(
        json.dumps([asdict(c) for c in cuentas], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def importar_desde_excel(rfc: str, xlsx_path: str) -> tuple[bool, str, list[CuentaContable]]:
    """
    Importa el catálogo de cuentas desde un Excel exportado de COI u otro ERP.
    Detecta columnas por nombre (flexible ante diferentes exports).
    Retorna (éxito, mensaje, lista_cuentas).
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws = wb.active
    except Exception as e:
        return False, f"No se pudo abrir el archivo: {e}", []

    # Buscar fila de encabezados (primeras 5 filas)
    hdrs: list[str] = []
    hdr_row = 1
    for r in range(1, 6):
        row_vals = [str(ws.cell(row=r, column=c).value or "") for c in range(1, 12)]
        if any(v.strip() for v in row_vals):
            hdrs = row_vals
            hdr_row = r
            break

    if not hdrs:
        return False, "No se encontró fila de encabezados.", []

    i_cta  = _col_idx(hdrs, "CUENTA", "NO.", "NUM")
    i_nom  = _col_idx(hdrs, "NOMBRE", "DESCRIPCION", "DESCRIPCI")
    i_tipo = _col_idx(hdrs, "TIPO", "NATURALEZA")
    i_niv  = _col_idx(hdrs, "NIVEL", "NIV")
    i_mon  = _col_idx(hdrs, "MONEDA", "MON")

    if i_cta is None or i_nom is None:
        return False, (
            "No se encontraron las columnas 'Cuenta' y 'Nombre'.\n"
            "Verifica que el Excel sea una exportación de catálogo de cuentas COI."
        ), []

    cuentas: list[CuentaContable] = []
    for r in range(hdr_row + 1, ws.max_row + 1):
        cta = str(ws.cell(row=r, column=i_cta + 1).value or "").strip()
        if not cta:
            continue
        nom = str(ws.cell(row=r, column=i_nom + 1).value or "").strip()
        if not nom:
            continue

        tipo_raw = str(ws.cell(row=r, column=i_tipo + 1).value or "D") if i_tipo is not None else "D"
        tipo = "A" if "A" in tipo_raw.upper() else "D"

        nivel_raw = ws.cell(row=r, column=i_niv + 1).value if i_niv is not None else 1
        try:
            nivel = int(nivel_raw or 1)
        except (ValueError, TypeError):
            nivel = 1

        moneda = str(ws.cell(row=r, column=i_mon + 1).value or "MN").strip() if i_mon is not None else "MN"
        if not moneda:
            moneda = "MN"

        cuentas.append(CuentaContable(cuenta=cta, nombre=nom,
                                      tipo=tipo, nivel=nivel, moneda=moneda))

    if not cuentas:
        return False, "No se encontraron cuentas en el archivo.", []

    guardar_catalogo(rfc, cuentas)
    return True, f"Catálogo importado: {len(cuentas)} cuentas.", cuentas


def buscar_cuentas(cuentas: list[CuentaContable], termino: str) -> list[CuentaContable]:
    """Filtra cuentas por número o nombre (búsqueda insensible a acentos)."""
    t = _norm(termino)
    return [c for c in cuentas if t in _norm(c.cuenta) or t in _norm(c.nombre)]


def cuenta_por_numero(cuentas: list[CuentaContable], numero: str) -> Optional[CuentaContable]:
    """Retorna la cuenta con ese número exacto, o None."""
    n = numero.strip()
    return next((c for c in cuentas if c.cuenta == n), None)
