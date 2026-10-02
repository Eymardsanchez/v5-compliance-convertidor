# -*- coding: utf-8 -*-
"""
reportes.py — Generador de reportes de Obligaciones Fiscales
=============================================================
Un solo exportador para las 4 pestañas del módulo (Panel General, Obligaciones,
Contribuciones/Saldos y Archivos), en los 4 formatos pedidos:

    xlsx · pdf · csv · txt

La pestaña arma un ``Reporte`` (título, subtítulo, encabezados, filas y metadatos
de empresa/periodo) y llama a :func:`exportar`. El contenido siempre corresponde a
la EMPRESA y al MES/AÑO seleccionados en esa pestaña.

PDF usa reportlab (ya instalado). Si falta, se avisa y se sugiere otro formato.
"""

from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass, field
from datetime import datetime


FORMATOS = ("xlsx", "pdf", "csv", "txt")

EXT_FORMATO = {"xlsx": ".xlsx", "pdf": ".pdf", "csv": ".csv", "txt": ".txt"}

NOMBRE_FORMATO = {
    "xlsx": "Excel (.xlsx)",
    "pdf":  "PDF (.pdf)",
    "csv":  "CSV (.csv)",
    "txt":  "Texto (.txt)",
}


@dataclass
class Reporte:
    titulo:      str = ""
    subtitulo:   str = ""
    empresa:     str = ""
    rfc:         str = ""
    periodo:     str = ""            # "Julio 2026"
    encabezados: list = field(default_factory=list)
    filas:       list = field(default_factory=list)   # list[list]
    resumen:     list = field(default_factory=list)   # [(etiqueta, valor), ...]

    def meta_lineas(self) -> list:
        out = []
        if self.empresa or self.rfc:
            out.append(f"Empresa: {self.empresa}  ({self.rfc})".strip())
        if self.periodo:
            out.append(f"Periodo: {self.periodo}")
        out.append(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
        return out


def nombre_sugerido(rep: Reporte, formato: str) -> str:
    """Nombre de archivo sugerido: Reporte_<titulo>_<RFC>_<periodo>.<ext>"""
    def _limpia(t):
        return "".join(c if (c.isalnum() or c in " -_") else "_" for c in (t or "")).strip().replace(" ", "_")
    partes = ["Reporte", _limpia(rep.titulo)]
    if rep.rfc:
        partes.append(_limpia(rep.rfc))
    if rep.periodo:
        partes.append(_limpia(rep.periodo))
    return "_".join([p for p in partes if p]) + EXT_FORMATO.get(formato, ".txt")


# ── Exportadores ─────────────────────────────────────────────────────────────
def _exportar_xlsx(rep: Reporte, ruta: str) -> str:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    thin = Border(*[Side(style="thin")] * 4)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (rep.titulo or "Reporte")[:31]
    ncols = max(1, len(rep.encabezados))

    r = 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    c = ws.cell(row=r, column=1, value=rep.titulo)
    c.font = Font(bold=True, size=14, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E79")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[r].height = 26
    r += 1
    if rep.subtitulo:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        c = ws.cell(row=r, column=1, value=rep.subtitulo)
        c.font = Font(italic=True, size=10, color="444444")
        c.alignment = Alignment(horizontal="center")
        r += 1
    for linea in rep.meta_lineas():
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        c = ws.cell(row=r, column=1, value=linea)
        c.font = Font(size=9, color="555555")
        r += 1
    r += 1

    if rep.encabezados:
        for j, h in enumerate(rep.encabezados, 1):
            c = ws.cell(row=r, column=j, value=h)
            c.font = Font(bold=True, color="FFFFFF", size=10)
            c.fill = PatternFill("solid", fgColor="2E75B6")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = thin
        ws.row_dimensions[r].height = 22
        hdr = r
        r += 1
        for fila in rep.filas:
            for j, v in enumerate(fila, 1):
                c = ws.cell(row=r, column=j, value=v)
                c.border = thin
                c.alignment = Alignment(vertical="center", wrap_text=isinstance(v, str) and len(str(v)) > 40)
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    c.number_format = '#,##0.00'
            r += 1
        ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
        ws.auto_filter.ref = (f"A{hdr}:{get_column_letter(ncols)}{max(hdr, r - 1)}")
        anchos = [0] * ncols
        for j, h in enumerate(rep.encabezados):
            anchos[j] = len(str(h))
        for fila in rep.filas:
            for j, v in enumerate(fila[:ncols]):
                anchos[j] = max(anchos[j], min(60, len(str(v if v is not None else ""))))
        for j, a in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(j)].width = max(10, a + 3)

    if rep.resumen:
        r += 1
        for etq, val in rep.resumen:
            ws.cell(row=r, column=1, value=str(etq)).font = Font(bold=True, size=10)
            ws.cell(row=r, column=2, value=val)
            r += 1

    wb.save(ruta)
    return ruta


def _exportar_csv(rep: Reporte, ruta: str) -> str:
    with io.open(ruta, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow([rep.titulo])
        if rep.subtitulo:
            w.writerow([rep.subtitulo])
        for linea in rep.meta_lineas():
            w.writerow([linea])
        w.writerow([])
        if rep.encabezados:
            w.writerow(rep.encabezados)
        for fila in rep.filas:
            w.writerow(["" if v is None else v for v in fila])
        if rep.resumen:
            w.writerow([])
            for etq, val in rep.resumen:
                w.writerow([etq, val])
    return ruta


def _exportar_txt(rep: Reporte, ruta: str) -> str:
    cols = len(rep.encabezados)
    anchos = [len(str(h)) for h in rep.encabezados]
    for fila in rep.filas:
        for j, v in enumerate(fila[:cols]):
            anchos[j] = max(anchos[j], min(50, len(str(v if v is not None else ""))))

    def _linea(vals):
        out = []
        for j, v in enumerate(vals[:cols]):
            out.append(str(v if v is not None else "").ljust(anchos[j])[:anchos[j]])
        return " | ".join(out)

    ancho_total = sum(anchos) + 3 * max(0, cols - 1)
    L = []
    L.append("=" * max(40, ancho_total))
    L.append(rep.titulo)
    if rep.subtitulo:
        L.append(rep.subtitulo)
    L.extend(rep.meta_lineas())
    L.append("=" * max(40, ancho_total))
    if rep.encabezados:
        L.append(_linea(rep.encabezados))
        L.append("-" * max(40, ancho_total))
        for fila in rep.filas:
            L.append(_linea(fila))
    if rep.resumen:
        L.append("-" * max(40, ancho_total))
        for etq, val in rep.resumen:
            L.append(f"{etq}: {val}")
    io.open(ruta, "w", encoding="utf-8").write("\n".join(L) + "\n")
    return ruta


def _exportar_pdf(rep: Reporte, ruta: str) -> str:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle)

    ancho_pag = landscape(letter) if len(rep.encabezados) > 5 else letter
    doc = SimpleDocTemplate(ruta, pagesize=ancho_pag,
                            leftMargin=1.4 * cm, rightMargin=1.4 * cm,
                            topMargin=1.4 * cm, bottomMargin=1.4 * cm,
                            title=rep.titulo)
    ss = getSampleStyleSheet()
    st_tit = ParagraphStyle("tit", parent=ss["Title"], fontSize=15, spaceAfter=2)
    st_sub = ParagraphStyle("sub", parent=ss["Normal"], fontSize=9,
                            textColor=colors.HexColor("#444444"), spaceAfter=1)
    st_cel = ParagraphStyle("cel", parent=ss["Normal"], fontSize=7.5, leading=9)
    st_hdr = ParagraphStyle("hdr", parent=ss["Normal"], fontSize=8, leading=10,
                            textColor=colors.white)

    flow = [Paragraph(rep.titulo, st_tit)]
    if rep.subtitulo:
        flow.append(Paragraph(rep.subtitulo, st_sub))
    for linea in rep.meta_lineas():
        flow.append(Paragraph(linea, st_sub))
    flow.append(Spacer(1, 8))

    if rep.encabezados:
        data = [[Paragraph(str(h), st_hdr) for h in rep.encabezados]]
        for fila in rep.filas:
            data.append([Paragraph(str(v if v is not None else ""), st_cel)
                         for v in fila[:len(rep.encabezados)]])
        t = Table(data, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E75B6")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9BB7D4")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F2F6FB")]),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        flow.append(t)

    if rep.resumen:
        flow.append(Spacer(1, 10))
        for etq, val in rep.resumen:
            flow.append(Paragraph(f"<b>{etq}:</b> {val}", st_cel))

    doc.build(flow)
    return ruta


def exportar(rep: Reporte, ruta: str, formato: str) -> str:
    """Escribe el reporte en el formato pedido. Devuelve la ruta generada."""
    formato = (formato or "").lower().lstrip(".")
    if formato == "xlsx":
        return _exportar_xlsx(rep, ruta)
    if formato == "csv":
        return _exportar_csv(rep, ruta)
    if formato == "txt":
        return _exportar_txt(rep, ruta)
    if formato == "pdf":
        try:
            return _exportar_pdf(rep, ruta)
        except ImportError as e:
            raise RuntimeError(
                "Para PDF se necesita la librería 'reportlab' "
                "(pip install reportlab). Usa Excel, CSV o TXT mientras tanto."
            ) from e
    raise ValueError(f"Formato no soportado: {formato}")
